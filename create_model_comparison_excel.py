import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
import pandas as pd

# File paths
xlsx_path = "model_comparison.xlsx"
csv_path = "model_comparison.csv"

# 1. Raw Data from the user's image
data = [
    {"Model": "Random Forest", "Accuracy": 0.931864, "Precision": 0.940054, "Recall": 0.922460, "F1": 0.931174, "AUC": 0.978839},
    {"Model": "KNN", "Accuracy": 0.921844, "Precision": 0.948791, "Recall": 0.891711, "F1": 0.919366, "AUC": 0.972913},
    {"Model": "Logistic Regression", "Accuracy": 0.789579, "Precision": 0.784494, "Recall": 0.798128, "F1": 0.791252, "AUC": 0.880748},
    {"Model": "Gaussian Naive Bayes", "Accuracy": 0.737475, "Precision": 0.720497, "Recall": 0.775401, "F1": 0.746941, "AUC": 0.796036},
]

# Export to CSV
df = pd.DataFrame(data)
df.to_csv(csv_path, index=False)
print(f"Generated CSV: {csv_path}")

# Create Excel Workbook with openpyxl
wb = openpyxl.Workbook()

# Sheet 1: Exact Image Replica
ws1 = wb.active
ws1.title = "Model Comparison"
ws1.views.sheetView[0].showGridLines = True

# Title row
ws1["A1"] = "4. Model Comparison"
ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")

# Headers
headers = ["Model", "Accuracy", "Precision", "Recall", "F1", "AUC"]
for col_idx, h in enumerate(headers, start=1):
    cell = ws1.cell(row=3, column=col_idx, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

# Thin border style
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Populate rows
for row_idx, row_data in enumerate(data, start=4):
    bg_color = "F9FBFD" if row_idx % 2 == 0 else "FFFFFF"
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    
    # Model name
    c_model = ws1.cell(row=row_idx, column=1, value=row_data["Model"])
    c_model.font = Font(name="Calibri", size=11, bold=(row_data["Model"] == "Random Forest"))
    c_model.alignment = Alignment(horizontal="left", vertical="center")
    c_model.fill = fill
    c_model.border = thin_border
    
    # Numeric values
    metrics = ["Accuracy", "Precision", "Recall", "F1", "AUC"]
    for col_idx, m in enumerate(metrics, start=2):
        cell = ws1.cell(row=row_idx, column=col_idx, value=row_data[m])
        cell.font = Font(name="Calibri", size=11)
        cell.number_format = "0.000000"
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.fill = fill
        cell.border = thin_border

# Auto-fit column widths
for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 5, 14)
ws1.column_dimensions["A"].width = 24

# Add Clustered Column Chart for Visual Representation
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Model Performance Benchmark Comparison"
chart.y_axis.title = "Score (0 - 1.0)"
chart.x_axis.title = "Metric"
chart.height = 12
chart.width = 18

chart_data = Reference(ws1, min_col=2, min_row=3, max_col=6, max_row=7)
chart_cats = Reference(ws1, min_col=1, min_row=4, max_row=7)
chart.add_data(chart_data, titles_from_data=True, from_rows=True)
chart.set_categories(Reference(ws1, min_col=2, min_row=3, max_col=6, max_row=3))
ws1.add_chart(chart, "A10")

# Sheet 2: Formatted Percentage view (ideal for slides and presentations)
ws2 = wb.create_sheet(title="Percentage View")
ws2.views.sheetView[0].showGridLines = True

ws2["A1"] = "Model Comparison (Percentage %)"
ws2["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")

for col_idx, h in enumerate(headers, start=1):
    cell = ws2.cell(row=3, column=col_idx, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

for row_idx, row_data in enumerate(data, start=4):
    bg_color = "F2F4F8" if row_idx % 2 == 0 else "FFFFFF"
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    
    c_model = ws2.cell(row=row_idx, column=1, value=row_data["Model"])
    c_model.font = Font(name="Calibri", size=11, bold=(row_data["Model"] == "Random Forest"))
    c_model.fill = fill
    c_model.border = thin_border
    
    for col_idx, m in enumerate(metrics, start=2):
        cell = ws2.cell(row=row_idx, column=col_idx, value=row_data[m])
        cell.font = Font(name="Calibri", size=11)
        cell.number_format = "0.00%"
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.fill = fill
        cell.border = thin_border

for col in ws2.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 5, 14)
ws2.column_dimensions["A"].width = 24

wb.save(xlsx_path)
print(f"Generated Excel Workbook: {xlsx_path}")
