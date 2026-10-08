import os
import pandas as pd
import numpy as np
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.ensemble import RandomForestClassifier
from sklearn.neighbors import KNeighborsClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.naive_bayes import GaussianNB
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, roc_auc_score
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

dataset_path = "unified_ev_fleet_dataset_final.xlsx"
if not os.path.exists(dataset_path):
    print("Dataset not found, skipping EV project benchmark.")
    exit(0)

print(f"Loading {dataset_path}...")
df = pd.read_excel(dataset_path)

# Target: trip_feasible (Binary classification: 1 vs 0)
target_col = "trip_feasible"
candidate_features = [
    'battery_capacity_kwh', 'max_range_km', 'motor_power_kw', 'vehicle_weight_kg',
    'state_of_charge_soc', 'state_of_health_soh', 'current_speed_kmph',
    'energy_consumed_kwh', 'trip_distance_km', 'trip_duration_min',
    'elevation_gain_m', 'outside_temperature_c', 'failure_risk_score'
]

features = [f for f in candidate_features if f in df.columns]
df_clean = df[features + [target_col]].dropna()

X = df_clean[features]
y = df_clean[target_col]

X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.25, random_state=42, stratify=y)

scaler = StandardScaler()
X_train_scaled = scaler.fit_transform(X_train)
X_test_scaled = scaler.transform(X_test)

models = {
    "Random Forest": RandomForestClassifier(n_estimators=100, max_depth=10, random_state=42),
    "KNN": KNeighborsClassifier(n_neighbors=5),
    "Logistic Regression": LogisticRegression(max_iter=1000, random_state=42),
    "Gaussian Naive Bayes": GaussianNB()
}

results = []
for name, clf in models.items():
    if name in ["KNN", "Logistic Regression", "Gaussian Naive Bayes"]:
        clf.fit(X_train_scaled, y_train)
        y_pred = clf.predict(X_test_scaled)
        y_prob = clf.predict_proba(X_test_scaled)[:, 1] if hasattr(clf, "predict_proba") else y_pred
    else:
        clf.fit(X_train, y_train)
        y_pred = clf.predict(X_test)
        y_prob = clf.predict_proba(X_test)[:, 1] if hasattr(clf, "predict_proba") else y_pred
        
    acc = accuracy_score(y_test, y_pred)
    prec = precision_score(y_test, y_pred, zero_division=0)
    rec = recall_score(y_test, y_pred, zero_division=0)
    f1 = f1_score(y_test, y_pred, zero_division=0)
    auc = roc_auc_score(y_test, y_prob)
    
    results.append({
        "Model": name,
        "Accuracy": acc,
        "Precision": prec,
        "Recall": rec,
        "F1": f1,
        "AUC": auc
    })

res_df = pd.DataFrame(results)
print("\n=== Real Dataset Benchmark Results (EV Fleet Trip Feasibility) ===")
print(res_df.to_string(index=False))

# Append to model_comparison.xlsx
xlsx_path = "model_comparison.xlsx"
wb = openpyxl.load_workbook(xlsx_path)

if "EV Fleet Project Results" in wb.sheetnames:
    del wb["EV Fleet Project Results"]

ws = wb.create_sheet(title="EV Fleet Project Results")
ws.views.sheetView[0].showGridLines = True

ws["A1"] = "4. Model Comparison (Trained on VoltFleet EV Dataset)"
ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")

headers = ["Model", "Accuracy", "Precision", "Recall", "F1", "AUC"]
for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=3, column=col_idx, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

metrics = ["Accuracy", "Precision", "Recall", "F1", "AUC"]
for row_idx, r in enumerate(results, start=4):
    bg_color = "F9FBFD" if row_idx % 2 == 0 else "FFFFFF"
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    
    c_m = ws.cell(row=row_idx, column=1, value=r["Model"])
    c_m.font = Font(name="Calibri", size=11, bold=(r["Model"] == "Random Forest"))
    c_m.fill = fill
    c_m.border = thin_border
    
    for col_idx, m in enumerate(metrics, start=2):
        cell = ws.cell(row=row_idx, column=col_idx, value=r[m])
        cell.font = Font(name="Calibri", size=11)
        cell.number_format = "0.000000"
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.fill = fill
        cell.border = thin_border

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 14)
ws.column_dimensions["A"].width = 24

wb.save(xlsx_path)
print(f"\nAppended real project benchmark to {xlsx_path}!")
