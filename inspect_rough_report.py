import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'Project PPT\Rough Report\VOLTFLEET AI (AutoRecovered).docx'
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- SAMPLE HEADINGS AND PARAGRAPHS ---")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt and (len(txt) < 80 or any(h in txt.upper() for h in ['CHAPTER', 'ABSTRACT', 'OBJECTIVE', 'LITERATURE', 'EXISTING', 'PROPOSED', 'ARCHITECTURE', 'MODULE', 'RESULT', 'REFERENCE', 'TABLE'])):
        print(f"[{idx}] {txt}")
