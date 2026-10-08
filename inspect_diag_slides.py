import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'Project PPT\project phase-1.pptx')

for idx in [7, 10]:
    slide = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        t = shape.text_frame.text.strip().replace('\n', ' ')[:60] if shape.has_text_frame else ('PICTURE' if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE else f'OTHER ({shape.shape_type})')
        print(f"  Shape {s_idx} ({shape.name}): {t} | pos: ({shape.left}, {shape.top}, {shape.width}, {shape.height})")
