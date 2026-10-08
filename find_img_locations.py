import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'Project PPT\Rough Report\VOLTFLEET AI (AutoRecovered).docx')

for idx, p in enumerate(doc.paragraphs):
    if "drawing" in p._element.xml:
        nearby = " // ".join([doc.paragraphs[i].text.strip() for i in range(max(0, idx-2), min(len(doc.paragraphs), idx+3)) if doc.paragraphs[i].text.strip()])
        print(f"Para {idx}: {nearby[:150]}")
