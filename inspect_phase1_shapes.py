import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'Project PPT\project phase-1.pptx')

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n=================== SLIDE {idx} ===================")
    for s_idx, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = " // ".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
        elif shape.has_table:
            txt = f"TABLE {len(shape.table.rows)}x{len(shape.table.columns)}"
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            txt = "PICTURE"
        else:
            txt = f"TYPE {shape.shape_type}"
        print(f"Shape {s_idx} ({shape.name}): '{txt[:80]}' | pos: ({shape.left}, {shape.top}, {shape.width}, {shape.height})")
