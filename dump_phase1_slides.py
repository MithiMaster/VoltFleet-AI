import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'Project PPT\project phase-1.pptx')

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n=================== SLIDE {idx} ===================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            print(f"--- Shape {s_idx} ({shape.name}) ---")
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f"  {txt}")
        elif shape.has_table:
            print(f"--- Shape {s_idx} (TABLE: {len(shape.table.rows)}x{len(shape.table.columns)}) ---")
            for r in list(shape.table.rows):
                print("  | " + " | ".join(c.text.strip().replace('\n', ' ') for c in r.cells))
