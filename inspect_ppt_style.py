import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Project PPT/VoltFleet_AI_EV_Fleet_Management_First_Review_PPT.pptx')
for idx in [0, 1, 5, 7]:
    slide = prs.slides[idx]
    print(f'=== SLIDE {idx+1} ===')
    for s in slide.shapes:
        if s.has_text_frame:
            for pi, p0 in enumerate(s.text_frame.paragraphs[:3]):
                txt = p0.text.strip()[:50]
                if not txt: continue
                r0 = p0.runs[0] if p0.runs else None
                fname = r0.font.name if r0 else 'None'
                fsize = r0.font.size.pt if r0 and r0.font.size else 'None'
                fcolor = r0.font.color.rgb if r0 and r0.font.color and r0.font.color.type == 1 else 'None'
                fbold = r0.font.bold if r0 else 'None'
                print(f'  [{s.name} p{pi}]: text="{txt}" | font={fname}, size={fsize}, bold={fbold}, color={fcolor}')
