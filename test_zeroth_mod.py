import pptx
import sys
import copy

sys.stdout.reconfigure(encoding='utf-8')

# Open the template
prs = pptx.Presentation('Project PPT/Zeroth Review.pptx')

print("Loaded Zeroth Review.pptx successfully!")
print(f"Number of slides: {len(prs.slides)}")

# Let's inspect slide 1 shapes and their texts
slide1 = prs.slides[0]
for idx, shape in enumerate(slide1.shapes):
    if shape.has_text_frame:
        print(f"Slide 1 Shape {idx} ({shape.name}):")
        for p in shape.text_frame.paragraphs:
            print(f"   P: '{p.text}'")

