import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Project PPT/Zeroth Review.pptx')
for idx in range(len(prs.slides)):
    slide = prs.slides[idx]
    print(f'=== ZEROTH REVIEW SLIDE {idx+1} ===')
    for s in slide.shapes:
        t = s.text_frame.text.strip().replace('\n', ' ')[:50] if s.has_text_frame else ('TABLE' if s.has_table else f'{s.shape_type}')
        print(f'  {s.name} ({s.shape_type}) pos:({s.left},{s.top},{s.width},{s.height}) text="{t}"')
