import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import sys

sys.stdout.reconfigure(encoding='utf-8')

src_path = 'Project PPT/Zeroth Review.pptx'
dest_path = 'Project PPT/VoltFleet_AI_EV_Fleet_Management_Zeroth_Review_PPT.pptx'

prs = pptx.Presentation(src_path)

def set_para_text(p, text, font_name="Inter", font_size=Pt(11), bold=False, color=RGBColor(0, 0, 0)):
    p.text = text
    p.font.name = font_name
    p.font.size = font_size
    p.font.bold = bold
    p.font.color.rgb = color

# ==========================================
# SLIDE 1: Title Slide
# ==========================================
slide1 = prs.slides[0]

# Shape 1 (Text 1): Project Title
sh_title = slide1.shapes[1]
tf = sh_title.text_frame
tf.word_wrap = True
tf.paragraphs[0].text = "VoltFleet AI: Intelligent EV Fleet Management, Predictive Battery Analytics & Route Optimization"
tf.paragraphs[0].font.name = "Inter"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126) # 00457E

# Shape 2 (Text 2): Review Subtitle
sh_sub = slide1.shapes[2]
tf_sub = sh_sub.text_frame
tf_sub.paragraphs[0].text = "Zeroth PROJECT REVIEW – Domain Presentation"
tf_sub.paragraphs[0].font.name = "Inter"
tf_sub.paragraphs[0].font.size = Pt(13)
tf_sub.paragraphs[0].font.bold = True
tf_sub.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180) # 4682B4

# Shape 4 (Text 3): Team Members
sh_team = slide1.shapes[4]
tf_team = sh_team.text_frame
tf_team.word_wrap = True
tf_team.clear()

p = tf_team.paragraphs[0]
set_para_text(p, "Team Members:", font_size=Pt(13), bold=True, color=RGBColor(0, 69, 126))

team_list = [
    "1. MITHILESH S     -  23UCS100",
    "2. MATHIVANAN M    -  23UCS099",
    "3. VISHNU K N      -  23UCS079",
    "4. NISHAN S        -  23UCS0108"
]
for member in team_list:
    p = tf_team.add_paragraph()
    set_para_text(p, member, font_size=Pt(11.5), bold=False, color=RGBColor(20, 20, 20))

# Shape 5 (Text 4): Project Guide
sh_guide = slide1.shapes[5]
tf_guide = sh_guide.text_frame
tf_guide.clear()
p = tf_guide.paragraphs[0]
set_para_text(p, "Guide Name: Mrs. P. NIRMALADEVI", font_size=Pt(12), bold=True, color=RGBColor(0, 69, 126))
p = tf_guide.add_paragraph()
set_para_text(p, "Associate Professor", font_size=Pt(11), bold=False, color=RGBColor(50, 50, 50))

# Shape 6 (Text 5): Department
sh_dept = slide1.shapes[6]
tf_dept = sh_dept.text_frame
p = tf_dept.paragraphs[0]
set_para_text(p, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", font_size=Pt(10.5), bold=True, color=RGBColor(0, 69, 126))

print("Updated Slide 1: Title & Team Details")

# ==========================================
# SLIDE 2: Domain Introduction
# ==========================================
slide2 = prs.slides[1]

# Shape 4 (Text 4): Bottom left tag
sh_tag2 = slide2.shapes[4]
sh_tag2.text_frame.paragraphs[0].text = "Intelligent EV Fleet Management"
sh_tag2.text_frame.paragraphs[0].font.name = "Inter"
sh_tag2.text_frame.paragraphs[0].font.size = Pt(8.5)
sh_tag2.text_frame.paragraphs[0].font.bold = True
sh_tag2.text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

# Shape 3 (Text 3): Domain descriptions
sh_doms = slide2.shapes[3]
tf_doms = sh_doms.text_frame
tf_doms.word_wrap = True
tf_doms.clear()

domain_sections = [
    ("Electric Vehicle (EV) Telematics & IoT Ingestion", " is the technology infrastructure that extracts real-time vehicular parameters (battery SOC, pack voltage, current velocity, motor power draw, inverter temperature) directly from CAN-bus interfaces. It is used to provide continuous, millisecond-level telemetry streaming to central cloud dispatchers."),
    ("Physics-Informed Machine Learning (PIML)", " is an emerging AI discipline that combines empirical data-driven learning with fundamental physical & electrochemical laws (aerodynamic drag, rolling resistance, gravitational potential, battery degradation). It is used to achieve highly accurate, generalizable driving range estimation and battery State of Health (SOH) forecasting across all operational duty cycles."),
    ("Geospatial Route Optimization & Topographical Profiling", " enables intelligent commercial fleet navigation by merging digital elevation models (DEM) with live road networks. It is used to calculate elevation power penalties, downhill regenerative braking energy recovery, and automated insertion of DC fast-charging stops to eliminate mid-transit stranding."),
    ("Driver Telematics Analytics & Behavioral AI", " utilizes machine learning to evaluate operator safety and energy efficiency. It is used to translate high-frequency driving patterns (acceleration jerk, harsh braking, speed consistency) into objective performance scores and personalized eco-driving guidance."),
    ("Cloud, Web & Fleet Database Systems", " provide the scalable operational backbone. They are used to ingest, securely store, and visualize mission-critical fleet telemetry through high-performance real-time interactive command centers.")
]

for idx, (head, body) in enumerate(domain_sections):
    p = tf_doms.paragraphs[0] if idx == 0 else tf_doms.add_paragraph()
    p.space_after = Pt(4)
    run_h = p.add_run()
    run_h.text = head
    run_h.font.name = "Inter"
    run_h.font.size = Pt(9.5)
    run_h.font.bold = True
    run_h.font.color.rgb = RGBColor(0, 69, 126)
    
    run_b = p.add_run()
    run_b.text = body
    run_b.font.name = "Inter"
    run_b.font.size = Pt(8.8)
    run_b.font.color.rgb = RGBColor(30, 30, 30)

print("Updated Slide 2: Domain Introduction")

# ==========================================
# SLIDE 3: Core Domain 1 - EV Telematics & CAN-Bus Ingestion
# ==========================================
slide3 = prs.slides[2]

# Shape 2 (Text 2): Title
slide3.shapes[2].text_frame.paragraphs[0].text = "EV Telematics & CAN-Bus Ingestion"
slide3.shapes[2].text_frame.paragraphs[0].font.name = "Inter"
slide3.shapes[2].text_frame.paragraphs[0].font.size = Pt(17)
slide3.shapes[2].text_frame.paragraphs[0].font.bold = True
slide3.shapes[2].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 3 (Text 3): Purpose
slide3.shapes[3].text_frame.clear()
p = slide3.shapes[3].text_frame.paragraphs[0]
r_p = p.add_run()
r_p.text = "Purpose: "
r_p.font.name = "Inter"
r_p.font.size = Pt(11)
r_p.font.bold = True
r_p.font.color.rgb = RGBColor(0, 69, 126)
r_b = p.add_run()
r_b.text = "Ingest, sanitize, and stream high-frequency vehicular telematics and sensor parameters into a unified central command hub."
r_b.font.name = "Inter"
r_b.font.size = Pt(10.5)

# Shape 4 (Text 4): Header "Our Idea"
slide3.shapes[4].text_frame.paragraphs[0].text = "Our Idea"
slide3.shapes[4].text_frame.paragraphs[0].font.name = "Inter"
slide3.shapes[4].text_frame.paragraphs[0].font.size = Pt(11)
slide3.shapes[4].text_frame.paragraphs[0].font.bold = True
slide3.shapes[4].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

# Shape 5 (Text 5): Techniques & Role
tf3_5 = slide3.shapes[5].text_frame
tf3_5.clear()

def add_bold_head_para(tf, head_txt, size=Pt(10.5), color=RGBColor(0, 69, 126)):
    p = tf.add_paragraph() if tf.paragraphs and tf.paragraphs[0].text else tf.paragraphs[0]
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = head_txt
    r.font.name = "Inter"
    r.font.size = size
    r.font.bold = True
    r.font.color.rgb = color
    return p

def add_body_para(tf, body_txt, size=Pt(9.2), color=RGBColor(40, 40, 40)):
    p = tf.add_paragraph()
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = body_txt
    r.font.name = "Inter"
    r.font.size = size
    r.font.color.rgb = color
    return p

add_bold_head_para(tf3_5, "Technologies / Techniques :", size=Pt(10.5), color=RGBColor(70, 130, 180))
for tech in [
    "1) CAN-bus OBD-II Protocol for high-frequency telemetry extraction.",
    "2) Multi-parameter live streaming: SOC %, pack voltage, velocity, motor power, inverter temp, GPS.",
    "3) Asynchronous non-blocking message processing with FastAPI & WebSockets.",
    "4) Time-series telemetry storage in SQLite / PostgreSQL with SQLAlchemy ORM.",
    "5) Edge buffering with automatic reconnection queue for zero data loss in remote transit zones."
]:
    add_body_para(tf3_5, tech, size=Pt(9.2))

add_bold_head_para(tf3_5, "Role in Our System", size=Pt(10.5), color=RGBColor(0, 69, 126))
add_body_para(tf3_5, "The Telematics Ingestion module serves as the primary real-time perception layer of VoltFleet AI, providing continuous ground-truth telemetry to the ML prediction engines and live command map.", size=Pt(9.2))

p_io = tf3_5.add_paragraph()
p_io.space_before = Pt(3)
r = p_io.add_run(); r.text = "Input: "; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 69, 126)
r = p_io.add_run(); r.text = "CAN-bus telemetry stream (GPS, speed, SOC, voltage, temp)    "; r.font.size = Pt(9.2)
r = p_io.add_run(); r.text = "Output: "; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 69, 126)
r = p_io.add_run(); r.text = "Sanitized time-series telemetry events and live fleet state"; r.font.size = Pt(9.2)

print("Updated Slide 3: EV Telematics & CAN-Bus Ingestion")

# ==========================================
# SLIDE 4: Core Domain 2 - Physics-Informed ML & Battery Range
# ==========================================
slide4 = prs.slides[3]

# Shape 1 (Text 1): Title
slide4.shapes[1].text_frame.paragraphs[0].text = "Physics-Informed ML & Battery Analytics"
slide4.shapes[1].text_frame.paragraphs[0].font.name = "Inter"
slide4.shapes[1].text_frame.paragraphs[0].font.size = Pt(17)
slide4.shapes[1].text_frame.paragraphs[0].font.bold = True
slide4.shapes[1].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 3 (Text 3): Bottom left tag
slide4.shapes[3].text_frame.paragraphs[0].text = "VoltFleet AI : Battery Intelligence Unit"
slide4.shapes[3].text_frame.paragraphs[0].font.name = "Inter"
slide4.shapes[3].text_frame.paragraphs[0].font.size = Pt(8.5)
slide4.shapes[3].text_frame.paragraphs[0].font.bold = True
slide4.shapes[3].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

# Shape 2 (Text 2): Content
tf4_2 = slide4.shapes[2].text_frame
tf4_2.clear()

# Purpose
p = tf4_2.paragraphs[0]
r = p.add_run(); r.text = "Purpose: "; r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0, 69, 126)
r = p.add_run(); r.text = "Accurately predict real-world EV driving range under dynamic operating conditions and model battery State of Health (SOH) degradation."; r.font.size = Pt(9.5)

add_bold_head_para(tf4_2, "Processing Pipeline", size=Pt(10.5), color=RGBColor(70, 130, 180))
p_pipe = tf4_2.add_paragraph()
r = p_pipe.add_run()
r.text = "Telematics Ingestion → Feature Engineering (SOC, Velocity, Temp, Elevation) → HistGradientBoosting Regressors → Range Deficit Warning → SOH Fade Tracking"
r.font.bold = True; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(20, 60, 100)

add_bold_head_para(tf4_2, "Techniques Used", size=Pt(10.5), color=RGBColor(70, 130, 180))
for tech in [
    "• Multi-variable gradient boosting regression (HistGradientBoosting / RandomForest)",
    "• Physical energy balance modeling (aerodynamic drag, rolling resistance, elevation potential energy)",
    "• Ambient temperature thermal efficiency compensation curves (-10°C to 50°C)",
    "• Cell delta-voltage and internal impedance monitoring for SOH capacity fade",
    "• Real-time range deficit warning algorithms before vehicle departure"
]:
    add_body_para(tf4_2, tech, size=Pt(8.8))

add_bold_head_para(tf4_2, "Role in Our System", size=Pt(10.5), color=RGBColor(0, 69, 126))
add_body_para(tf4_2, "Acts as the core battery intelligence engine, replacing naive dashboard averages with physically grounded range forecasts to completely eliminate range anxiety.", size=Pt(8.8))

p_io = tf4_2.add_paragraph()
p_io.space_before = Pt(2)
r = p_io.add_run(); r.text = "Input: "; r.font.bold = True; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(0, 69, 126)
r = p_io.add_run(); r.text = "Battery capacity (kWh), SOC %, velocity, outside temp, elevation, HVAC    "; r.font.size = Pt(8.8)
r = p_io.add_run(); r.text = "Output: "; r.font.bold = True; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(0, 69, 126)
r = p_io.add_run(); r.text = "Dynamic remaining range (km), range deficit flag, battery SOH %"; r.font.size = Pt(8.8)

print("Updated Slide 4: Physics-Informed ML & Battery Analytics")

# ==========================================
# SLIDE 5: Core Domain 3 - Geospatial Route Optimization
# ==========================================
slide5 = prs.slides[4]

# Shape 1 (Text 1): Title
slide5.shapes[1].text_frame.paragraphs[0].text = "Geospatial Routing & Elevation Profiling"
slide5.shapes[1].text_frame.paragraphs[0].font.name = "Inter"
slide5.shapes[1].text_frame.paragraphs[0].font.size = Pt(17)
slide5.shapes[1].text_frame.paragraphs[0].font.bold = True
slide5.shapes[1].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 3 (Text 3): Bottom left tag
slide5.shapes[3].text_frame.paragraphs[0].text = "VoltFleet AI : Navigation & Logistics Brain"
slide5.shapes[3].text_frame.paragraphs[0].font.name = "Inter"
slide5.shapes[3].text_frame.paragraphs[0].font.size = Pt(8.5)
slide5.shapes[3].text_frame.paragraphs[0].font.bold = True
slide5.shapes[3].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

# Shape 2 (Text 2): Content
tf5_2 = slide5.shapes[2].text_frame
tf5_2.clear()

p = tf5_2.paragraphs[0]
r = p.add_run(); r.text = "Purpose: "; r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0, 69, 126)
r = p.add_run(); r.text = "Formulate energy-optimal multi-stop transit routes balancing distance, topographical elevation gradients, and automated charging waypoint insertion."; r.font.size = Pt(9.5)

add_bold_head_para(tf5_2, "Major Processing Tasks", size=Pt(10.5), color=RGBColor(70, 130, 180))
for task in [
    "• Multi-stop origin, destination, and intermediate waypoint route planning",
    "• 3D Digital Elevation Modeling (DEM) integration to extract terrain altitude profiles",
    "• Gravitational energy penalty calculation on uphill climbs",
    "• Downhill regenerative braking recovery estimation",
    "• Battery discharge curve simulation along the transit corridor",
    "• Automated DC fast-charging stop recommendation when SOC projected < 20%",
    "• Total transit duration and cost estimation under Time-of-Use (ToU) electricity tariffs"
]:
    add_body_para(tf5_2, task, size=Pt(8.8))

add_bold_head_para(tf5_2, "Role in Our System", size=Pt(10.5), color=RGBColor(0, 69, 126))
add_body_para(tf5_2, "Acts as the operational navigational brain of VoltFleet AI, guaranteeing zero en-route vehicle stranding across cross-city transit missions.", size=Pt(8.8))

p_io = tf5_2.add_paragraph()
p_io.space_before = Pt(2)
r = p_io.add_run(); r.text = "Input: "; r.font.bold = True; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(0, 69, 126)
r = p_io.add_run(); r.text = "Origin/destination coordinates, intermediate stops, initial battery SOC, vehicle specs    "; r.font.size = Pt(8.8)
r = p_io.add_run(); r.text = "Output: "; r.font.bold = True; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(0, 69, 126)
r = p_io.add_run(); r.text = "Energy-optimized route polyline, elevation profile, recommended charging waypoints"; r.font.size = Pt(8.8)

print("Updated Slide 5: Geospatial Routing & Elevation Profiling")

# ==========================================
# SLIDE 6: Sub Domains and Issues (Part 1)
# ==========================================
slide6 = prs.slides[5]

# Shape 2 (Text 2): Col 1 Header
slide6.shapes[2].text_frame.paragraphs[0].text = "EV Telematics & Real-Time Ingestion"
slide6.shapes[2].text_frame.paragraphs[0].font.name = "Inter"
slide6.shapes[2].text_frame.paragraphs[0].font.size = Pt(11)
slide6.shapes[2].text_frame.paragraphs[0].font.bold = True
slide6.shapes[2].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 3 (Text 3): Col 1 Bullets
tf6_3 = slide6.shapes[3].text_frame
tf6_3.clear()
for issue in [
    "• Sensor jitter, packet loss, and delayed telemetry synchronization over cellular networks.",
    "• Variable CAN-bus baud rates and proprietary manufacturer message definitions across diverse EV brands.",
    "• High computational overhead of handling continuous multi-vehicle streaming data.",
    "• Managing offline telemetry buffering when vehicles traverse dead zones or tunnels."
]:
    p = tf6_3.add_paragraph() if tf6_3.paragraphs[0].text else tf6_3.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = issue; r.font.name = "Inter"; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(30, 30, 30)

# Shape 5 (Text 5): Col 2 Header
slide6.shapes[5].text_frame.paragraphs[0].text = "Battery Health (SOH) & Thermal Dynamics"
slide6.shapes[5].text_frame.paragraphs[0].font.name = "Inter"
slide6.shapes[5].text_frame.paragraphs[0].font.size = Pt(11)
slide6.shapes[5].text_frame.paragraphs[0].font.bold = True
slide6.shapes[5].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 6 (Text 6): Col 2 Bullets
tf6_6 = slide6.shapes[6].text_frame
tf6_6.clear()
for issue in [
    "• Non-linear lithium-ion electrochemical degradation affected by ambient temperature and C-rate.",
    "• Thermal runaway risks during rapid DC fast charging in extreme summer climates.",
    "• Inaccurate onboard BMS state-of-charge drift over repeated partial charging cycles.",
    "• Complex cell-level voltage imbalances requiring continuous diagnostic filtering."
]:
    p = tf6_6.add_paragraph() if tf6_6.paragraphs[0].text else tf6_6.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = issue; r.font.name = "Inter"; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(30, 30, 30)

# Shape 8 (Text 8): Col 3 Header
slide6.shapes[8].text_frame.paragraphs[0].text = "Topographical Range & Energy Modeling"
slide6.shapes[8].text_frame.paragraphs[0].font.name = "Inter"
slide6.shapes[8].text_frame.paragraphs[0].font.size = Pt(11)
slide6.shapes[8].text_frame.paragraphs[0].font.bold = True
slide6.shapes[8].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 9 (Text 9): Col 3 Bullets
tf6_9 = slide6.shapes[9].text_frame
tf6_9.clear()
for issue in [
    "• High variability in real-world energy consumption due to aggressive headwind and elevation shifts.",
    "• Inaccurate static vehicle dashboard range estimators that disregard elevation and HVAC load.",
    "• Difficulty modeling regenerative braking recovery efficiency under differing battery SOC levels.",
    "• Lack of standardized road grade datasets for precise hill climb energy computation."
]:
    p = tf6_9.add_paragraph() if tf6_9.paragraphs[0].text else tf6_9.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = issue; r.font.name = "Inter"; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(30, 30, 30)

print("Updated Slide 6: Sub Domains and Issues (Part 1)")

# ==========================================
# SLIDE 7: Sub Domains and Issues (Part 2)
# ==========================================
slide7 = prs.slides[6]

# Shape 2 (Text 2): Col 1 Header
slide7.shapes[2].text_frame.paragraphs[0].text = "Driver Telematics & Behavioral Scoring"
slide7.shapes[2].text_frame.paragraphs[0].font.name = "Inter"
slide7.shapes[2].text_frame.paragraphs[0].font.size = Pt(11)
slide7.shapes[2].text_frame.paragraphs[0].font.bold = True
slide7.shapes[2].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 3 (Text 3): Col 1 Bullets
tf7_3 = slide7.shapes[3].text_frame
tf7_3.clear()
for issue in [
    "• High-frequency throttle jerk and aggressive braking accelerating mechanical wear and energy drain.",
    "• Fairly isolating driver behavior from external road congestion, weather, and payload weights.",
    "• Providing transparent, explainable coaching metrics without causing driver dispute or distrust.",
    "• Translating raw CAN-bus accelerations into standardized safety and efficiency scores."
]:
    p = tf7_3.add_paragraph() if tf7_3.paragraphs[0].text else tf7_3.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = issue; r.font.name = "Inter"; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(30, 30, 30)

# Shape 5 (Text 5): Col 2 Header
slide7.shapes[5].text_frame.paragraphs[0].text = "Multi-Stop Route Planning & Charging Logistics"
slide7.shapes[5].text_frame.paragraphs[0].font.name = "Inter"
slide7.shapes[5].text_frame.paragraphs[0].font.size = Pt(11)
slide7.shapes[5].text_frame.paragraphs[0].font.bold = True
slide7.shapes[5].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 6 (Text 6): Col 2 Bullets
tf7_6 = slide7.shapes[6].text_frame
tf7_6.clear()
for issue in [
    "• Dynamic charging station availability and unpredictable queue waiting times at public hubs.",
    "• Incompatible charging connector standards (CCS2, GB/T, CHAdeMO) across commercial vehicle models.",
    "• Balancing trade-offs between shortest-distance routes and elevation-efficient bypass corridors.",
    "• Incorporating variable Time-of-Use (ToU) electricity tariffs into cost-effective transit stops."
]:
    p = tf7_6.add_paragraph() if tf7_6.paragraphs[0].text else tf7_6.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = issue; r.font.name = "Inter"; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(30, 30, 30)

# Shape 8 (Text 8): Col 3 Header
slide7.shapes[8].text_frame.paragraphs[0].text = "Edge-to-Cloud Telemetry Synchronization"
slide7.shapes[8].text_frame.paragraphs[0].font.name = "Inter"
slide7.shapes[8].text_frame.paragraphs[0].font.size = Pt(11)
slide7.shapes[8].text_frame.paragraphs[0].font.bold = True
slide7.shapes[8].text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

# Shape 9 (Text 9): Col 3 Bullets
tf7_9 = slide7.shapes[9].text_frame
tf7_9.clear()
for issue in [
    "• Ensuring zero-trust cryptographic verification of vehicle telematics packets against spoofing.",
    "• Maintaining low-latency bidirectional communication between dispatchers and remote drivers.",
    "• Scalable data storage architecture for gigabytes of longitudinal vehicle sensor logs.",
    "• Granular role-based access control protecting commercial fleet financial and operational data."
]:
    p = tf7_9.add_paragraph() if tf7_9.paragraphs[0].text else tf7_9.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run(); r.text = issue; r.font.name = "Inter"; r.font.size = Pt(8.8); r.font.color.rgb = RGBColor(30, 30, 30)

print("Updated Slide 7: Sub Domains and Issues (Part 2)")

# ==========================================
# SLIDE 8: Proposed System Modules
# ==========================================
slide8 = prs.slides[7]

# Module 1 (Top Left)
slide8.shapes[3].text_frame.paragraphs[0].text = "Telematics Ingestion & Fleet Command Center"
slide8.shapes[3].text_frame.paragraphs[0].font.name = "Inter"
slide8.shapes[3].text_frame.paragraphs[0].font.size = Pt(11)
slide8.shapes[3].text_frame.paragraphs[0].font.bold = True
slide8.shapes[3].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

tf8_4 = slide8.shapes[4].text_frame
tf8_4.clear()
for b in [
    "• Ingests CAN-bus telemetry streams in real time.",
    "• Renders interactive geospatial Leaflet command map.",
    "• Visualizes vehicle status pins, live battery SOC %, speed, and health diagnostics."
]:
    p = tf8_4.add_paragraph() if tf8_4.paragraphs[0].text else tf8_4.paragraphs[0]
    p.space_after = Pt(3)
    r = p.add_run(); r.text = b; r.font.name = "Inter"; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(30, 30, 30)

# Module 2 (Bottom Left)
slide8.shapes[6].text_frame.paragraphs[0].text = "Physics-Informed ML Range & Battery Health"
slide8.shapes[6].text_frame.paragraphs[0].font.name = "Inter"
slide8.shapes[6].text_frame.paragraphs[0].font.size = Pt(11)
slide8.shapes[6].text_frame.paragraphs[0].font.bold = True
slide8.shapes[6].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

tf8_7 = slide8.shapes[7].text_frame
tf8_7.clear()
for b in [
    "• Evaluates HistGradientBoosting regression on velocity, SOC, temp, and terrain.",
    "• Tracks battery SOH capacity fade across charge cycles.",
    "• Triggers proactive thermal runaway and cell-level overheat alerts."
]:
    p = tf8_7.add_paragraph() if tf8_7.paragraphs[0].text else tf8_7.paragraphs[0]
    p.space_after = Pt(3)
    r = p.add_run(); r.text = b; r.font.name = "Inter"; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(30, 30, 30)

# Module 3 (Top Right)
slide8.shapes[9].text_frame.paragraphs[0].text = "Dynamic Multi-Stop Route & Charging Optimizer"
slide8.shapes[9].text_frame.paragraphs[0].font.name = "Inter"
slide8.shapes[9].text_frame.paragraphs[0].font.size = Pt(11)
slide8.shapes[9].text_frame.paragraphs[0].font.bold = True
slide8.shapes[9].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

tf8_10 = slide8.shapes[10].text_frame
tf8_10.clear()
for b in [
    "• Calculates 3D altitude profiles and climbing power penalties.",
    "• Models downhill regenerative braking energy recovery.",
    "• Automates DC fast-charging stop recommendations to eliminate range anxiety."
]:
    p = tf8_10.add_paragraph() if tf8_10.paragraphs[0].text else tf8_10.paragraphs[0]
    p.space_after = Pt(3)
    r = p.add_run(); r.text = b; r.font.name = "Inter"; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(30, 30, 30)

# Module 4 (Bottom Right)
slide8.shapes[12].text_frame.paragraphs[0].text = "ML Driver Performance & Predictive Maintenance"
slide8.shapes[12].text_frame.paragraphs[0].font.name = "Inter"
slide8.shapes[12].text_frame.paragraphs[0].font.size = Pt(11)
slide8.shapes[12].text_frame.paragraphs[0].font.bold = True
slide8.shapes[12].text_frame.paragraphs[0].font.color.rgb = RGBColor(70, 130, 180)

tf8_13 = slide8.shapes[13].text_frame
tf8_13.clear()
for b in [
    "• Multiple Linear Regression evaluates Safety, Efficiency, Battery Care, and Route Compliance.",
    "• Delivers automated, personalized coaching to operators.",
    "• Auto-generates predictive service tickets based on sensor wear patterns."
]:
    p = tf8_13.add_paragraph() if tf8_13.paragraphs[0].text else tf8_13.paragraphs[0]
    p.space_after = Pt(3)
    r = p.add_run(); r.text = b; r.font.name = "Inter"; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(30, 30, 30)

print("Updated Slide 8: Proposed System Modules")

# ==========================================
# SLIDE 9: Thank You
# ==========================================
slide9 = prs.slides[8]

# Add Project Title and Details on Slide 9
# Let's inspect slide 9 shapes:
# Shape 1 is "Thank You"
# Let's add subtitle and guide info in Shape 1 or new textbox
tf9 = slide9.shapes[1].text_frame
tf9.paragraphs[0].text = "Thank You"
tf9.paragraphs[0].font.name = "Inter"
tf9.paragraphs[0].font.size = Pt(32)
tf9.paragraphs[0].font.bold = True
tf9.paragraphs[0].font.color.rgb = RGBColor(0, 69, 126)

p_sub = tf9.add_paragraph()
p_sub.space_before = Pt(14)
r = p_sub.add_run()
r.text = "VoltFleet AI: Intelligent EV Fleet Management, Predictive Battery Analytics & Route Optimization"
r.font.name = "Inter"
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = RGBColor(70, 130, 180)

p_qa = tf9.add_paragraph()
p_qa.space_before = Pt(8)
r = p_qa.add_run()
r.text = "Questions & Discussion"
r.font.name = "Inter"
r.font.size = Pt(12)
r.font.bold = False
r.font.color.rgb = RGBColor(50, 50, 50)

p_team = tf9.add_paragraph()
p_team.space_before = Pt(12)
r = p_team.add_run()
r.text = "Team: Mithilesh S (23UCS100) | Mathivanan M (23UCS099) | Vishnu K N (23UCS079) | Nishan S (23UCS0108)"
r.font.name = "Inter"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(90, 90, 90)

p_guide = tf9.add_paragraph()
p_guide.space_before = Pt(4)
r = p_guide.add_run()
r.text = "Project Guide: Mrs. P. NIRMALADEVI, Associate Professor | Department of Computer Science & Engineering"
r.font.name = "Inter"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(90, 90, 90)

print("Updated Slide 9: Thank You & Credits")

# Save Presentation
prs.save(dest_path)
print(f"\nSuccessfully generated Zeroth Review PPT at: {dest_path}")
