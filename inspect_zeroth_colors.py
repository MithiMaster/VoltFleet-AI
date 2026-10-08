import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Project PPT/Zeroth Review.pptx')
for idx in range(len(prs.slides)):
    slide = prs.slides[idx]
    print(f'=== SLIDE {idx+1} ===')
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    f = r.font
                    color = f.color.rgb if f.color and f.color.type == 1 else ('theme' if f.color else 'None')
                    print(f'  [{s.name}] "{r.text[:30]}" | font={f.name}, sz={f.size.pt if f.size else "-"}, c={color}, b={f.bold}')
