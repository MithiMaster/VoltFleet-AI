import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

template_path = r'Project PPT\project phase-1.pptx'
prs = pptx.Presentation(template_path)

print(f"=== TEMPLATE: {template_path} ===")
print(f"Slide count: {len(prs.slides)}, Width: {prs.slide_width}, Height: {prs.slide_height}")

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n--- SLIDE {idx} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            texts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if texts:
                sample = " // ".join(texts)
                print(f"  Shape {s_idx} ({shape.name}): {sample[:140]}")
        elif shape.has_table:
            print(f"  Shape {s_idx} (TABLE): {len(shape.table.rows)} rows x {len(shape.table.columns)} cols")
            for r in list(shape.table.rows)[:3]:
                print("    | " + " | ".join(c.text.strip().replace('\n', ' ') for c in r.cells))
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  Shape {s_idx} (PICTURE): {shape.name}")
