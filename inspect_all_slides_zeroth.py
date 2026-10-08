import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Project PPT/Zeroth Review.pptx')

for s_idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {s_idx + 1} ===================")
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            full_txt = " // ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if full_txt:
                print(f"Shape {sh_idx} ({shape.name}): '{full_txt[:100]}'")

