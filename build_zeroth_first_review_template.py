import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

# Source template to clone styles, masters and layouts
template_ppt = r'Project PPT/VoltFleet_AI_EV_Fleet_Management_First_Review_PPT.pptx'
output_ppt = r'Project PPT/VoltFleet_AI_EV_Fleet_Management_Zeroth_Review_PPT.pptx'
header_img_path = r'Project PPT/college_header.png'

prs = pptx.Presentation(template_ppt)
print(f"Loaded template: {template_ppt} ({len(prs.slides)} slides)")
print(f"Dimensions: {prs.slide_width} x {prs.slide_height}")

# Common Colors
COLOR_BG_GRAY = RGBColor(242, 242, 242)     # #F2F2F2
COLOR_BG_WHITE = RGBColor(255, 255, 255)    # #FFFFFF
COLOR_BLACK = RGBColor(0, 0, 0)             # #000000
COLOR_DARK = RGBColor(26, 26, 26)           # #1A1A1A
COLOR_GRAY_TEXT = RGBColor(60, 60, 60)      # #3C3C3C
COLOR_MUTED = RGBColor(100, 100, 100)       # #646464
COLOR_CARD_BORDER = RGBColor(204, 204, 204) # #CCCCCC
COLOR_ACCENT_BLUE = RGBColor(0, 69, 126)    # #00457E
COLOR_FOOTER_BG = RGBColor(235, 238, 242)   # #EBEEF2

# Geometry Constants (matching First Review template)
LEFT_MARGIN = 457200
CONTENT_WIDTH = 11274552
TITLE_TOP = 201168
TITLE_HEIGHT = 645160
LINE_TOP = 824357
LINE_HEIGHT = 27432
FOOTER_TOP = 6187515
FOOTER_HEIGHT = 330000

def set_slide_bg(slide, rgb_color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb_color

def set_run(p, text, font_name="Times New Roman", font_size=Pt(16), bold=False, italic=False, color=COLOR_DARK):
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    # Also set on paragraph level as fallback
    p.font.name = font_name
    p.font.size = font_size
    p.font.bold = bold
    p.font.italic = italic
    if color:
        p.font.color.rgb = color
    return run

def add_header_line(slide):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        LEFT_MARGIN, LINE_TOP, CONTENT_WIDTH, LINE_HEIGHT
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_DARK
    line.line.fill.solid()
    line.line.color.rgb = COLOR_DARK
    line.line.width = Pt(0.75)
    return line

def add_slide_title(slide, title_text):
    tx_box = slide.shapes.add_textbox(LEFT_MARGIN, TITLE_TOP, CONTENT_WIDTH, TITLE_HEIGHT)
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p, title_text, font_name="Times New Roman", font_size=Pt(34), bold=True, color=COLOR_BLACK)
    add_header_line(slide)
    return tx_box

def add_card(slide, left, top, width, height, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.fill.solid()
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def add_footer_io(slide, io_text):
    card = add_card(slide, LEFT_MARGIN, FOOTER_TOP, CONTENT_WIDTH, FOOTER_HEIGHT, bg_color=COLOR_FOOTER_BG, border_color=COLOR_CARD_BORDER)
    tx_box = slide.shapes.add_textbox(LEFT_MARGIN + 91440, FOOTER_TOP + 35000, CONTENT_WIDTH - 182880, FOOTER_HEIGHT - 70000)
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    set_run(p, io_text, font_name="Times New Roman", font_size=Pt(13.5), bold=True, color=COLOR_DARK)
    return card

# Remove all slides except slide 1
print("Clearing slides from template...")
for i in range(len(prs.slides) - 1, 0, -1):
    rId = prs.slides._sldIdLst[i].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[i]

blank_layout = prs.slide_layouts[6]

# ==============================================================================
# SLIDE 1: Title Slide (Update existing slide 1)
# ==============================================================================
print("Formatting Slide 1 (Title Slide)...")
slide1 = prs.slides[0]
set_slide_bg(slide1, COLOR_BG_WHITE)

# Update texts with precise Times New Roman runs
for shape in slide1.shapes:
    if shape.has_text_frame:
        tf = shape.text_frame
        full_text = " ".join([p.text.strip() for p in tf.paragraphs if p.text.strip()])
        
        if "IV Year" in full_text or "Final Year Project" in full_text or "Zeroth" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "Zeroth Project Review - Domain Presentation", font_name="Times New Roman", font_size=Pt(15), bold=True, color=COLOR_BLACK)
        elif "Assistant Professor" in full_text or "Associate Professor" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "Associate Professor", font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)
        elif "TEAM MEMBERS:" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "TEAM MEMBERS:", font_name="Times New Roman", font_size=Pt(15), bold=True, color=COLOR_BLACK)
        elif "PROJECT GUIDE:" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "PROJECT GUIDE:", font_name="Times New Roman", font_size=Pt(15), bold=True, color=COLOR_BLACK)
        elif "MITHILESH" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "MITHILESH S  -  23UCS100", font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)
        elif "MATHIVANAN" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "MATHIVANAN M -  23UCS099", font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)
        elif "VISHNU" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "VISHNU K N -  23UCS079", font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)
        elif "NISHAN" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "NISHAN S - 23UCS0108", font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)
        elif "NIRMALADEVI" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "Mrs. P. NIRMALADEVI", font_name="Times New Roman", font_size=Pt(14), bold=True, color=COLOR_BLACK)
        elif "Department Of Computer Science" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p, "Department Of Computer Science And Engineering", font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)
        elif "VoltFleet AI:" in full_text:
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_run(p, "VoltFleet AI: Intelligent EV Fleet Management, Predictive Battery Analytics & Route Optimization", font_name="Times New Roman", font_size=Pt(28), bold=True, color=COLOR_BLACK)

# ==============================================================================
# SLIDE 2: Domain Introduction
# ==============================================================================
print("Building Slide 2 (Domain Introduction)...")
slide2 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide2, COLOR_BG_GRAY)
add_slide_title(slide2, "Domain Introduction")

# Subtitle / context tag
tag_box = slide2.shapes.add_textbox(LEFT_MARGIN, 860000, CONTENT_WIDTH, 260000)
tf_tag = tag_box.text_frame
p_tag = tf_tag.paragraphs[0]
p_tag.alignment = PP_ALIGN.CENTER
set_run(p_tag, "Core Technological Pillars of VoltFleet AI & Modern EV Telematics", font_name="Times New Roman", font_size=Pt(15), bold=False, italic=True, color=COLOR_MUTED)

# Content box with the 5 domains
tb2 = slide2.shapes.add_textbox(LEFT_MARGIN, 1200000, CONTENT_WIDTH, 5200000)
tf2 = tb2.text_frame
tf2.word_wrap = True
tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0

domains = [
    ("Electric Vehicle (EV) Telematics & IoT Ingestion", 
     "is the technology infrastructure that extracts real-time vehicular parameters (battery SOC, pack voltage, current velocity, motor power draw, inverter temperature) directly from CAN-bus interfaces. It provides continuous, millisecond-level telemetry streaming to central cloud dispatchers."),
    
    ("Physics-Informed Machine Learning (PIML)", 
     "is an emerging AI discipline combining empirical data-driven learning with fundamental physical & electrochemical laws (aerodynamic drag, rolling resistance, gravitational potential, battery degradation) to achieve highly accurate, generalizable range estimation and battery State of Health (SOH) forecasting across all operational duty cycles."),
    
    ("Geospatial Route Optimization & Topographical Profiling", 
     "enables intelligent commercial fleet navigation by merging digital elevation models (DEM) with live road networks, calculating elevation power penalties, downhill regenerative braking energy recovery, and automated insertion of DC fast-charging stops to eliminate mid-transit stranding."),
    
    ("Driver Telematics Analytics & Behavioral AI", 
     "utilizes machine learning to evaluate operator safety and energy efficiency, translating high-frequency driving patterns (acceleration jerk, harsh braking, speed consistency) into objective performance scores and personalized eco-driving guidance."),
    
    ("Cloud, Web & Fleet Database Systems", 
     "provide the scalable operational backbone to ingest, securely store, and visualize mission-critical fleet telemetry through high-performance real-time interactive command centers.")
]

for idx, (title, desc) in enumerate(domains):
    p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
    p.space_after = Pt(11)
    p.space_before = Pt(3)
    
    set_run(p, "❖  ", font_name="Times New Roman", font_size=Pt(17), bold=True, color=COLOR_BLACK)
    set_run(p, f"{title}: ", font_name="Times New Roman", font_size=Pt(17), bold=True, color=COLOR_BLACK)
    set_run(p, desc, font_name="Times New Roman", font_size=Pt(16.5), bold=False, color=COLOR_DARK)

# ==============================================================================
# SLIDE 3: Domain Focus : EV Telematics & CAN-Bus Ingestion
# ==============================================================================
print("Building Slide 3 (EV Telematics & CAN-Bus Ingestion)...")
slide3 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide3, COLOR_BG_GRAY)
add_slide_title(slide3, "Domain Focus : EV Telematics & CAN-Bus Ingestion")

# Purpose Card
add_card(slide3, LEFT_MARGIN, 920000, CONTENT_WIDTH, 680000, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
tb_p = slide3.shapes.add_textbox(LEFT_MARGIN + 180000, 960000, CONTENT_WIDTH - 360000, 600000)
tf_p = tb_p.text_frame
tf_p.word_wrap = True
tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0
p = tf_p.paragraphs[0]
set_run(p, "Purpose: ", font_name="Times New Roman", font_size=Pt(18), bold=True, color=COLOR_BLACK)
set_run(p, "Ingest, sanitize, and stream high-frequency vehicular telematics and sensor parameters into a unified central command hub.", font_name="Times New Roman", font_size=Pt(18), bold=False, color=COLOR_DARK)

# Content Box (Tech + Role)
tb_c = slide3.shapes.add_textbox(LEFT_MARGIN, 1720000, CONTENT_WIDTH, 4350000)
tf_c = tb_c.text_frame
tf_c.word_wrap = True
tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0

# Technologies / Techniques Header
p = tf_c.paragraphs[0]
p.space_after = Pt(8)
set_run(p, "Technologies & Techniques:", font_name="Times New Roman", font_size=Pt(20), bold=True, color=COLOR_BLACK)

techs = [
    "CAN-bus OBD-II Protocol for high-frequency telemetry extraction (50–100 Hz).",
    "Multi-parameter live streaming: SOC %, pack voltage, velocity, motor power, inverter temp, GPS.",
    "Asynchronous non-blocking message processing with FastAPI & WebSockets.",
    "Time-series telemetry storage in SQLite / PostgreSQL with SQLAlchemy ORM.",
    "Edge buffering with automatic reconnection queue for zero data loss in remote transit zones."
]

for t in techs:
    p = tf_c.add_paragraph()
    p.space_after = Pt(6)
    set_run(p, f"•  {t}", font_name="Times New Roman", font_size=Pt(17.5), bold=False, color=COLOR_DARK)

# Role Header
p = tf_c.add_paragraph()
p.space_before = Pt(12)
p.space_after = Pt(6)
set_run(p, "Role in Our System:", font_name="Times New Roman", font_size=Pt(20), bold=True, color=COLOR_BLACK)

p = tf_c.add_paragraph()
set_run(p, "The Telematics Ingestion module serves as the primary real-time perception layer of VoltFleet AI, providing continuous ground-truth telemetry to the ML prediction engines and live command map.", font_name="Times New Roman", font_size=Pt(17.5), bold=False, color=COLOR_DARK)

# Footer Bar
add_footer_io(slide3, "⚙ Input: CAN-bus telemetry stream (GPS, speed, SOC, voltage, temp)  |  Output: Sanitized time-series telemetry events & live fleet state")

# ==============================================================================
# SLIDE 4: Domain Focus : Physics-Informed ML & Battery Analytics
# ==============================================================================
print("Building Slide 4 (Physics-Informed ML & Battery Analytics)...")
slide4 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide4, COLOR_BG_GRAY)
add_slide_title(slide4, "Domain Focus : Physics-Informed ML & Battery Analytics")

# Purpose Card
add_card(slide4, LEFT_MARGIN, 920000, CONTENT_WIDTH, 680000, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
tb_p = slide4.shapes.add_textbox(LEFT_MARGIN + 180000, 960000, CONTENT_WIDTH - 360000, 600000)
tf_p = tb_p.text_frame
tf_p.word_wrap = True
tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0
p = tf_p.paragraphs[0]
set_run(p, "Purpose: ", font_name="Times New Roman", font_size=Pt(18), bold=True, color=COLOR_BLACK)
set_run(p, "Accurately predict real-world EV driving range under dynamic operating conditions and model battery State of Health (SOH) degradation.", font_name="Times New Roman", font_size=Pt(18), bold=False, color=COLOR_DARK)

# Pipeline Card
add_card(slide4, LEFT_MARGIN, 1680000, CONTENT_WIDTH, 520000, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
tb_pipe = slide4.shapes.add_textbox(LEFT_MARGIN + 180000, 1720000, CONTENT_WIDTH - 360000, 440000)
tf_pipe = tb_pipe.text_frame
tf_pipe.word_wrap = True
tf_pipe.margin_left = tf_pipe.margin_right = tf_pipe.margin_top = tf_pipe.margin_bottom = 0
p = tf_pipe.paragraphs[0]
set_run(p, "Processing Pipeline:  ", font_name="Times New Roman", font_size=Pt(15.5), bold=True, color=COLOR_BLACK)
set_run(p, "Telematics Ingestion  ➔  Feature Engineering (SOC, Velocity, Temp, Elevation)  ➔  HistGradientBoosting  ➔  Range Deficit Alert  ➔  SOH Fade Tracking", font_name="Times New Roman", font_size=Pt(15), bold=False, color=COLOR_DARK)

# Content Box (Techniques + Role)
tb_c = slide4.shapes.add_textbox(LEFT_MARGIN, 2280000, CONTENT_WIDTH, 3800000)
tf_c = tb_c.text_frame
tf_c.word_wrap = True
tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0

# Techniques Header
p = tf_c.paragraphs[0]
p.space_after = Pt(6)
set_run(p, "Techniques Used:", font_name="Times New Roman", font_size=Pt(20), bold=True, color=COLOR_BLACK)

techs4 = [
    "Multi-variable gradient boosting regression (HistGradientBoosting / RandomForest) trained on vehicle dynamics.",
    "Physical energy balance modeling (aerodynamic drag, rolling resistance, elevation potential energy).",
    "Ambient temperature thermal efficiency compensation curves (-10°C to 50°C).",
    "Cell delta-voltage and internal impedance monitoring for SOH capacity fade.",
    "Real-time range deficit warning algorithms before vehicle departure."
]

for t in techs4:
    p = tf_c.add_paragraph()
    p.space_after = Pt(5)
    set_run(p, f"•  {t}", font_name="Times New Roman", font_size=Pt(17), bold=False, color=COLOR_DARK)

# Role Header
p = tf_c.add_paragraph()
p.space_before = Pt(10)
p.space_after = Pt(4)
set_run(p, "Role in Our System:", font_name="Times New Roman", font_size=Pt(20), bold=True, color=COLOR_BLACK)

p = tf_c.add_paragraph()
set_run(p, "Acts as the core battery intelligence engine, replacing naive dashboard averages with physically grounded range forecasts to completely eliminate range anxiety.", font_name="Times New Roman", font_size=Pt(17), bold=False, color=COLOR_DARK)

# Footer Bar
add_footer_io(slide4, "⚙ Input: Battery capacity (kWh), SOC %, velocity, outside temp, elevation, HVAC  |  Output: Dynamic remaining range (km), range deficit flag, battery SOH %")

# ==============================================================================
# SLIDE 5: Domain Focus : Geospatial Routing & Elevation Profiling
# ==============================================================================
print("Building Slide 5 (Geospatial Routing & Elevation Profiling)...")
slide5 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide5, COLOR_BG_GRAY)
add_slide_title(slide5, "Domain Focus : Geospatial Routing & Elevation Profiling")

# Purpose Card
add_card(slide5, LEFT_MARGIN, 920000, CONTENT_WIDTH, 680000, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
tb_p = slide5.shapes.add_textbox(LEFT_MARGIN + 180000, 960000, CONTENT_WIDTH - 360000, 600000)
tf_p = tb_p.text_frame
tf_p.word_wrap = True
tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0
p = tf_p.paragraphs[0]
set_run(p, "Purpose: ", font_name="Times New Roman", font_size=Pt(18), bold=True, color=COLOR_BLACK)
set_run(p, "Formulate energy-optimal multi-stop transit routes balancing distance, topographical elevation gradients, and automated charging waypoint insertion.", font_name="Times New Roman", font_size=Pt(18), bold=False, color=COLOR_DARK)

# Content Box (Tasks + Role)
tb_c = slide5.shapes.add_textbox(LEFT_MARGIN, 1700000, CONTENT_WIDTH, 4350000)
tf_c = tb_c.text_frame
tf_c.word_wrap = True
tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0

# Major Processing Tasks Header
p = tf_c.paragraphs[0]
p.space_after = Pt(6)
set_run(p, "Major Processing Tasks:", font_name="Times New Roman", font_size=Pt(20), bold=True, color=COLOR_BLACK)

tasks5 = [
    "Multi-stop origin, destination, and intermediate waypoint route planning via Leaflet & OpenStreetMap.",
    "3D Digital Elevation Modeling (DEM) integration to extract continuous terrain altitude profiles.",
    "Gravitational energy penalty calculation on uphill climbs and road slope variations.",
    "Downhill regenerative braking energy recovery estimation for dynamic battery re-charge.",
    "Battery discharge curve simulation along the transit corridor.",
    "Automated DC fast-charging stop recommendation when projected SOC drops below safe threshold (< 20%).",
    "Total transit duration and cost estimation factoring Time-of-Use (ToU) electricity tariffs."
]

for t in tasks5:
    p = tf_c.add_paragraph()
    p.space_after = Pt(4)
    set_run(p, f"•  {t}", font_name="Times New Roman", font_size=Pt(16.5), bold=False, color=COLOR_DARK)

# Role Header
p = tf_c.add_paragraph()
p.space_before = Pt(8)
p.space_after = Pt(4)
set_run(p, "Role in Our System:", font_name="Times New Roman", font_size=Pt(20), bold=True, color=COLOR_BLACK)

p = tf_c.add_paragraph()
set_run(p, "Acts as the operational navigational brain of VoltFleet AI, guaranteeing zero en-route vehicle stranding across cross-city transit missions.", font_name="Times New Roman", font_size=Pt(16.5), bold=False, color=COLOR_DARK)

# Footer Bar
add_footer_io(slide5, "⚙ Input: Origin/destination coords, intermediate stops, initial SOC, vehicle specs  |  Output: Energy-optimized route polyline, elevation profile, charging waypoints")

# ==============================================================================
# SLIDE 6: Sub Domains and Issues (Part 1)
# ==============================================================================
print("Building Slide 6 (Sub Domains and Issues - Part 1)...")
slide6 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide6, COLOR_BG_GRAY)
add_slide_title(slide6, "Sub Domains and Issues (Part 1)")

# Subtitle / Context
tag_box = slide6.shapes.add_textbox(LEFT_MARGIN, 860000, CONTENT_WIDTH, 260000)
tf_tag = tag_box.text_frame
p_tag = tf_tag.paragraphs[0]
p_tag.alignment = PP_ALIGN.CENTER
set_run(p_tag, "Critical Engineering & Research Challenges in Core EV Technologies", font_name="Times New Roman", font_size=Pt(15), bold=False, italic=True, color=COLOR_MUTED)

# 3 Cards Side-by-Side
card_w = 3550000
card_gap = 310000
card_top = 1180000
card_h = 5350000

cards_data_6 = [
    {
        "title": "EV Telematics & Real-Time Ingestion",
        "issues": [
            "Sensor jitter, packet loss, and delayed telemetry synchronization over cellular networks.",
            "Variable CAN-bus baud rates and proprietary manufacturer message definitions across diverse EV brands.",
            "High computational overhead of handling continuous multi-vehicle streaming data.",
            "Managing offline telemetry buffering when vehicles traverse dead zones or tunnels."
        ]
    },
    {
        "title": "Battery Health (SOH) & Thermal Dynamics",
        "issues": [
            "Non-linear lithium-ion electrochemical degradation affected by ambient temperature and C-rate.",
            "Thermal runaway risks during rapid DC fast charging in extreme summer climates.",
            "Inaccurate onboard BMS state-of-charge drift over repeated partial charging cycles.",
            "Complex cell-level voltage imbalances requiring continuous diagnostic filtering."
        ]
    },
    {
        "title": "Topographical Range & Energy Modeling",
        "issues": [
            "High variability in real-world energy consumption due to aggressive headwind and elevation shifts.",
            "Inaccurate static vehicle dashboard range estimators that disregard elevation and HVAC load.",
            "Difficulty modeling regenerative braking recovery efficiency under differing battery SOC levels.",
            "Lack of standardized road grade datasets for precise hill climb energy computation."
        ]
    }
]

for i, c_data in enumerate(cards_data_6):
    c_left = LEFT_MARGIN + i * (card_w + card_gap)
    add_card(slide6, c_left, card_top, card_w, card_h, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
    
    # Header box inside card
    hdr_box = slide6.shapes.add_textbox(c_left + 160000, card_top + 160000, card_w - 320000, 700000)
    tf_h = hdr_box.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = tf_h.margin_bottom = 0
    p = tf_h.paragraphs[0]
    set_run(p, c_data["title"], font_name="Times New Roman", font_size=Pt(18), bold=True, color=COLOR_BLACK)
    
    # Inner divider line
    div = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left + 160000, card_top + 880000, card_w - 320000, Pt(1.5))
    div.fill.solid()
    div.fill.fore_color.rgb = COLOR_DARK
    div.line.fill.solid()
    div.line.color.rgb = COLOR_DARK
    
    # Bullet text box
    txt_box = slide6.shapes.add_textbox(c_left + 160000, card_top + 980000, card_w - 320000, card_h - 1100000)
    tf_t = txt_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    
    for idx, iss in enumerate(c_data["issues"]):
        p = tf_t.paragraphs[0] if idx == 0 else tf_t.add_paragraph()
        p.space_after = Pt(12)
        set_run(p, f"•  {iss}", font_name="Times New Roman", font_size=Pt(15), bold=False, color=COLOR_DARK)

# ==============================================================================
# SLIDE 7: Sub Domains and Issues (Part 2)
# ==============================================================================
print("Building Slide 7 (Sub Domains and Issues - Part 2)...")
slide7 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide7, COLOR_BG_GRAY)
add_slide_title(slide7, "Sub Domains and Issues (Part 2)")

# Subtitle / Context
tag_box = slide7.shapes.add_textbox(LEFT_MARGIN, 860000, CONTENT_WIDTH, 260000)
tf_tag = tag_box.text_frame
p_tag = tf_tag.paragraphs[0]
p_tag.alignment = PP_ALIGN.CENTER
set_run(p_tag, "Operational, Logistical & Security Challenges in Commercial Fleet Deployment", font_name="Times New Roman", font_size=Pt(15), bold=False, italic=True, color=COLOR_MUTED)

cards_data_7 = [
    {
        "title": "Driver Telematics & Behavioral Scoring",
        "issues": [
            "High-frequency throttle jerk and aggressive braking accelerating mechanical wear and energy drain.",
            "Fairly isolating driver behavior from external road congestion, weather, and payload weights.",
            "Providing transparent, explainable coaching metrics without causing driver dispute or distrust.",
            "Translating raw CAN-bus accelerations into standardized safety and efficiency scores."
        ]
    },
    {
        "title": "Multi-Stop Route Planning & Charging Logistics",
        "issues": [
            "Dynamic charging station availability and unpredictable queue waiting times at public hubs.",
            "Incompatible charging connector standards (CCS2, GB/T, CHAdeMO) across commercial vehicle models.",
            "Balancing trade-offs between shortest-distance routes and elevation-efficient bypass corridors.",
            "Incorporating variable Time-of-Use (ToU) electricity tariffs into cost-effective transit stops."
        ]
    },
    {
        "title": "Edge-to-Cloud Telemetry Synchronization",
        "issues": [
            "Ensuring zero-trust cryptographic verification of vehicle telematics packets against spoofing.",
            "Maintaining low-latency bidirectional communication between dispatchers and remote drivers.",
            "Scalable data storage architecture for gigabytes of longitudinal vehicle sensor logs.",
            "Granular role-based access control protecting commercial fleet financial and operational data."
        ]
    }
]

for i, c_data in enumerate(cards_data_7):
    c_left = LEFT_MARGIN + i * (card_w + card_gap)
    add_card(slide7, c_left, card_top, card_w, card_h, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
    
    hdr_box = slide7.shapes.add_textbox(c_left + 160000, card_top + 160000, card_w - 320000, 700000)
    tf_h = hdr_box.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = tf_h.margin_bottom = 0
    p = tf_h.paragraphs[0]
    set_run(p, c_data["title"], font_name="Times New Roman", font_size=Pt(18), bold=True, color=COLOR_BLACK)
    
    div = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left + 160000, card_top + 880000, card_w - 320000, Pt(1.5))
    div.fill.solid()
    div.fill.fore_color.rgb = COLOR_DARK
    div.line.fill.solid()
    div.line.color.rgb = COLOR_DARK
    
    txt_box = slide7.shapes.add_textbox(c_left + 160000, card_top + 980000, card_w - 320000, card_h - 1100000)
    tf_t = txt_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    
    for idx, iss in enumerate(c_data["issues"]):
        p = tf_t.paragraphs[0] if idx == 0 else tf_t.add_paragraph()
        p.space_after = Pt(12)
        set_run(p, f"•  {iss}", font_name="Times New Roman", font_size=Pt(15), bold=False, color=COLOR_DARK)

# ==============================================================================
# SLIDE 8: System Modules
# ==============================================================================
print("Building Slide 8 (System Modules)...")
slide8 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide8, COLOR_BG_GRAY)
add_slide_title(slide8, "VoltFleet AI : System Modules")

# Subtitle / Context
tag_box = slide8.shapes.add_textbox(LEFT_MARGIN, 860000, CONTENT_WIDTH, 260000)
tf_tag = tag_box.text_frame
p_tag = tf_tag.paragraphs[0]
p_tag.alignment = PP_ALIGN.CENTER
set_run(p_tag, "Modular Architecture for Comprehensive Intelligent EV Fleet Management", font_name="Times New Roman", font_size=Pt(15), bold=False, italic=True, color=COLOR_MUTED)

# 2x2 Grid Layout
m_card_w = 5480000
m_card_h = 2500000
col1_left = LEFT_MARGIN
col2_left = LEFT_MARGIN + m_card_w + 314552
row1_top = 1180000
row2_top = 3900000

modules_data = [
    {
        "col": col1_left, "top": row1_top,
        "title": "Module 1 : Telematics Ingestion & Fleet Command Center",
        "bullets": [
            "Ingests CAN-bus telemetry streams in real time (GPS, velocity, SOC, voltage, motor power).",
            "Renders interactive geospatial Leaflet command map with live vehicle status pins.",
            "Visualizes fleet KPI metrics, battery SOC %, speed alerts, and health diagnostics."
        ]
    },
    {
        "col": col2_left, "top": row1_top,
        "title": "Module 2 : Physics-Informed ML Range & Battery Health",
        "bullets": [
            "Evaluates HistGradientBoosting regression on velocity, SOC, temp, and terrain elevation.",
            "Tracks battery SOH capacity fade and degradation across charge cycles.",
            "Triggers proactive thermal runaway warnings and cell-level overheat alerts."
        ]
    },
    {
        "col": col1_left, "top": row2_top,
        "title": "Module 3 : Dynamic Multi-Stop Route & Charging Optimizer",
        "bullets": [
            "Calculates 3D altitude profiles and climbing gravitational power penalties.",
            "Models downhill regenerative braking energy recovery and battery discharge curves.",
            "Automates DC fast-charging stop recommendations (< 20% SOC) to eliminate range anxiety."
        ]
    },
    {
        "col": col2_left, "top": row2_top,
        "title": "Module 4 : ML Driver Performance & Predictive Maintenance",
        "bullets": [
            "Multiple Linear Regression evaluates Safety, Efficiency, Battery Care, and Route Compliance.",
            "Delivers automated, personalized eco-driving coaching to fleet operators.",
            "Auto-generates predictive service tickets based on sensor wear patterns and diagnostic codes."
        ]
    }
]

for m in modules_data:
    add_card(slide8, m["col"], m["top"], m_card_w, m_card_h, bg_color=COLOR_BG_WHITE, border_color=COLOR_CARD_BORDER)
    
    # Title
    t_box = slide8.shapes.add_textbox(m["col"] + 160000, m["top"] + 140000, m_card_w - 320000, 480000)
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p = tf_t.paragraphs[0]
    set_run(p, m["title"], font_name="Times New Roman", font_size=Pt(17), bold=True, color=COLOR_BLACK)
    
    # Divider
    div = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, m["col"] + 160000, m["top"] + 660000, m_card_w - 320000, Pt(1))
    div.fill.solid()
    div.fill.fore_color.rgb = COLOR_DARK
    div.line.fill.solid()
    div.line.color.rgb = COLOR_DARK
    
    # Bullets
    b_box = slide8.shapes.add_textbox(m["col"] + 160000, m["top"] + 740000, m_card_w - 320000, m_card_h - 820000)
    tf_b = b_box.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
    for idx, b in enumerate(m["bullets"]):
        p = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
        p.space_after = Pt(8)
        set_run(p, f"•  {b}", font_name="Times New Roman", font_size=Pt(15.5), bold=False, color=COLOR_DARK)

# ==============================================================================
# SLIDE 9: Thank You Slide
# ==============================================================================
print("Building Slide 9 (Thank You)...")
slide9 = prs.slides.add_slide(blank_layout)
set_slide_bg(slide9, COLOR_BG_WHITE)

# Header Banner Image
slide9.shapes.add_picture(header_img_path, 457201, 27900, 11274552, 1285717)

# Upper divider line
l1 = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, 1371600, CONTENT_WIDTH, 36576)
l1.fill.solid()
l1.fill.fore_color.rgb = COLOR_BLACK
l1.line.fill.solid()
l1.line.color.rgb = COLOR_BLACK

# Title & Thank you box
tx_ty = slide9.shapes.add_textbox(LEFT_MARGIN, 1480000, CONTENT_WIDTH, 1550000)
tf_ty = tx_ty.text_frame
tf_ty.word_wrap = True
tf_ty.margin_left = tf_ty.margin_right = tf_ty.margin_top = tf_ty.margin_bottom = 0

p1 = tf_ty.paragraphs[0]
p1.alignment = PP_ALIGN.CENTER
p1.space_after = Pt(10)
set_run(p1, "THANK YOU", font_name="Times New Roman", font_size=Pt(38), bold=True, color=COLOR_BLACK)

p2 = tf_ty.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.space_after = Pt(8)
set_run(p2, "VoltFleet AI: Intelligent EV Fleet Management, Predictive Battery Analytics & Route Optimization", font_name="Times New Roman", font_size=Pt(18), bold=True, color=COLOR_DARK)

p3 = tf_ty.add_paragraph()
p3.alignment = PP_ALIGN.CENTER
set_run(p3, "Questions & Discussion", font_name="Times New Roman", font_size=Pt(16), bold=False, italic=True, color=COLOR_MUTED)

# Lower divider line
l2 = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, 3154680, CONTENT_WIDTH, 36576)
l2.fill.solid()
l2.fill.fore_color.rgb = COLOR_BLACK
l2.line.fill.solid()
l2.line.color.rgb = COLOR_BLACK

# Subtitle / Department line
tb_sub = slide9.shapes.add_textbox(548640, 3364992, 5486400, 323165)
p_sub = tb_sub.text_frame.paragraphs[0]
set_run(p_sub, "Zeroth Project Review - Domain Presentation", font_name="Times New Roman", font_size=Pt(15), bold=True, color=COLOR_BLACK)

# Left column: Team Members
tb_team_hdr = slide9.shapes.add_textbox(548640, 3767328, 5486400, 320040)
p = tb_team_hdr.text_frame.paragraphs[0]
set_run(p, "TEAM MEMBERS:", font_name="Times New Roman", font_size=Pt(15), bold=True, color=COLOR_BLACK)

members = [
    ("MITHILESH S  -  23UCS100", 4187952),
    ("MATHIVANAN M -  23UCS099", 4690872),
    ("VISHNU K N -  23UCS079", 5146167),
    ("NISHAN S - 23UCS0108", 5600700)
]

for name, top_pos in members:
    tb_m = slide9.shapes.add_textbox(548640, top_pos, 5486400, 306705)
    p = tb_m.text_frame.paragraphs[0]
    set_run(p, name, font_name="Times New Roman", font_size=Pt(14), bold=False, color=COLOR_BLACK)

# Right column: Project Guide
tb_guide_hdr = slide9.shapes.add_textbox(6611112, 3493008, 5029200, 347472)
p = tb_guide_hdr.text_frame.paragraphs[0]
set_run(p, "PROJECT GUIDE:", font_name="Times New Roman", font_size=Pt(15), bold=True, color=COLOR_BLACK)

guide_details = [
    ("Mrs. P. NIRMALADEVI", 3950208, True),
    ("Associate Professor", 4425696, False),
    ("Department Of Computer Science And Engineering", 4901184, False)
]

for g_txt, top_pos, is_b in guide_details:
    tb_g = slide9.shapes.add_textbox(6611112, top_pos, 5029200, 306705)
    p = tb_g.text_frame.paragraphs[0]
    set_run(p, g_txt, font_name="Times New Roman", font_size=Pt(14), bold=is_b, color=COLOR_BLACK)

prs.save(output_ppt)
print(f"\n=======================================================")
print(f"SUCCESS: Saved {len(prs.slides)} slides to {output_ppt}")
print(f"=======================================================")
