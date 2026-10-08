import os
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
import fitz

def generate_outcomes_pdf():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_pdf = os.path.join(script_dir, "VoltFleet_AI_Project_Outcomes_1_Page.pdf")

    # Target letter size: 612 x 792 pt
    # Matching Smart_Lift_Project_Outcomes_1_Page.pdf layout
    doc = SimpleDocTemplate(
        output_pdf,
        pagesize=letter,
        leftMargin=57.6,
        rightMargin=57.6,
        topMargin=46.0,
        bottomMargin=46.0
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=16,
        leading=21,
        alignment=TA_CENTER,
        spaceAfter=14
    )

    intro_style = ParagraphStyle(
        'IntroText',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=12,
        leading=15.8,
        alignment=TA_JUSTIFY,
        spaceAfter=11
    )

    bullet_style = ParagraphStyle(
        'BulletText',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=11.5,
        leading=14.8,
        alignment=TA_JUSTIFY,
        leftIndent=18,
        firstLineIndent=-18,
        spaceAfter=6.5
    )

    conclusion_style = ParagraphStyle(
        'ConclusionText',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=11.5,
        leading=15.2,
        alignment=TA_JUSTIFY,
        spaceBefore=8
    )

    story = []
    story.append(Paragraph("OUTCOMES", title_style))

    intro_text = (
        "The VoltFleet AI platform successfully achieved the core objectives of the intelligent EV fleet "
        "management software prototype stage. The project demonstrates the integration of Artificial Intelligence, "
        "Physics-Informed Machine Learning, IoT telematics ingestion, elevation-aware route optimization, "
        "and an interactive real-time dashboard for comprehensive commercial fleet operations."
    )
    story.append(Paragraph(intro_text, intro_style))

    bullets = [
        "Successfully ingested and processed multi-parameter EV telematics streams, including battery State of Charge (SOC), State of Health (SOH), pack voltage, motor power draw, cell temperature, and live GPS coordinates.",
        "Implemented a Physics-Informed Machine Learning (HistGradientBoosting) regression model to accurately forecast real-world remaining range by factoring battery SOC, vehicle velocity, ambient temperature (-10°C to 50°C), and auxiliary HVAC loads.",
        "Built a Dynamic Multi-Stop Route Optimizer using 3D digital elevation modeling (DEM) that accounts for gravitational power penalties on uphill ascents and regenerative braking energy recapture on downhill descents.",
        "Automated 150 kW DC fast-charging waypoint insertion along transit corridors, ensuring battery SOC remains above the critical 20% safety reserve buffer to eliminate mid-transit stranding.",
        "Developed a 4-Pillar AI Driver Performance Scoring engine using Multiple Linear Regression to quantify Safety, Efficiency, Battery Care, and Route Compliance with personalized coaching recommendations.",
        "Implemented an Off-Peak Charging Scheduler utilizing Time-of-Use (ToU) electricity tariffs to automate depot charging schedules, achieving up to 31.4% operational energy cost reduction.",
        "Established a sensor-driven Predictive Maintenance module tracking cumulative odometer milestones, cell delta-voltage anomalies, and thermal stress to auto-generate timely service work orders.",
        "Developed an interactive web-based Fleet Command Center featuring high-contrast night-mode Leaflet GIS mapping, dynamic fleet KPI aggregations, and individual vehicle health diagnostics.",
        "Integrated an interactive &ldquo;What-If&rdquo; Mission Simulator enabling fleet dispatchers to evaluate route feasibility, payload impacts, and weather conditions prior to vehicle dispatch.",
        "Implemented secure Role-Based Access Control (RBAC) with cryptographic bcrypt password hashing and JSON Web Tokens (JWT) to isolate Fleet Dispatcher management capabilities from Driver shift workflows."
    ]

    for b in bullets:
        story.append(Paragraph(f"&bull;&nbsp;&nbsp;{b}", bullet_style))

    concl_text = (
        "Overall, the completed prototype proves the core intelligent decision-making logic of the VoltFleet AI "
        "system, including physics-informed battery range estimation, topographical route optimization, driver "
        "behavior analytics, automated charging management, and interactive GIS visualization. Hardware CAN-bus "
        "transceiver integration and automated Vehicle-to-Grid (V2G) load balancing are planned as the next development phase."
    )
    story.append(Paragraph(concl_text, conclusion_style))

    doc.build(story)
    print(f"Successfully generated {output_pdf}")

    # Verify page count
    verify_doc = fitz.open(output_pdf)
    print(f"Total pages generated: {len(verify_doc)}")
    assert len(verify_doc) == 1, f"Expected 1 page, got {len(verify_doc)}"

def generate_outcomes_docx():
    import docx
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_docx = os.path.join(script_dir, "VoltFleet_AI_Project_Outcomes_1_Page.docx")

    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.7)
        sec.bottom_margin = Inches(0.7)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(14)
    r = p_title.add_run("OUTCOMES")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)

    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.space_after = Pt(10)
    r = p_intro.add_run(
        "The VoltFleet AI platform successfully achieved the core objectives of the intelligent EV fleet "
        "management software prototype stage. The project demonstrates the integration of Artificial Intelligence, "
        "Physics-Informed Machine Learning, IoT telematics ingestion, elevation-aware route optimization, "
        "and an interactive real-time dashboard for comprehensive commercial fleet operations."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    bullets = [
        "Successfully ingested and processed multi-parameter EV telematics streams, including battery State of Charge (SOC), State of Health (SOH), pack voltage, motor power draw, cell temperature, and live GPS coordinates.",
        "Implemented a Physics-Informed Machine Learning (HistGradientBoosting) regression model to accurately forecast real-world remaining range by factoring battery SOC, vehicle velocity, ambient temperature (-10°C to 50°C), and auxiliary HVAC loads.",
        "Built a Dynamic Multi-Stop Route Optimizer using 3D digital elevation modeling (DEM) that accounts for gravitational power penalties on uphill ascents and regenerative braking energy recapture on downhill descents.",
        "Automated 150 kW DC fast-charging waypoint insertion along transit corridors, ensuring battery SOC remains above the critical 20% safety reserve buffer to eliminate mid-transit stranding.",
        "Developed a 4-Pillar AI Driver Performance Scoring engine using Multiple Linear Regression to quantify Safety, Efficiency, Battery Care, and Route Compliance with personalized coaching recommendations.",
        "Implemented an Off-Peak Charging Scheduler utilizing Time-of-Use (ToU) electricity tariffs to automate depot charging schedules, achieving up to 31.4% operational energy cost reduction.",
        "Established a sensor-driven Predictive Maintenance module tracking cumulative odometer milestones, cell delta-voltage anomalies, and thermal stress to auto-generate timely service work orders.",
        "Developed an interactive web-based Fleet Command Center featuring high-contrast night-mode Leaflet GIS mapping, dynamic fleet KPI aggregations, and individual vehicle health diagnostics.",
        "Integrated an interactive \"What-If\" Mission Simulator enabling fleet dispatchers to evaluate route feasibility, payload impacts, and weather conditions prior to vehicle dispatch.",
        "Implemented secure Role-Based Access Control (RBAC) with cryptographic bcrypt password hashing and JSON Web Tokens (JWT) to isolate Fleet Dispatcher management capabilities from Driver shift workflows."
    ]

    for b in bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Inches(0.25)
        p_b.paragraph_format.space_after = Pt(6)
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_bullet = p_b.add_run("•  ")
        r_bullet.font.name = "Times New Roman"
        r_bullet.font.size = Pt(12)
        r_text = p_b.add_run(b)
        r_text.font.name = "Times New Roman"
        r_text.font.size = Pt(11.5)

    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.space_before = Pt(8)
    p_concl.paragraph_format.space_after = Pt(0)
    p_concl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_concl = p_concl.add_run(
        "Overall, the completed prototype proves the core intelligent decision-making logic of the VoltFleet AI "
        "system, including physics-informed battery range estimation, topographical route optimization, driver "
        "behavior analytics, automated charging management, and interactive GIS visualization. Hardware CAN-bus "
        "transceiver integration and automated Vehicle-to-Grid (V2G) load balancing are planned as the next development phase."
    )
    r_concl.bold = True
    r_concl.font.name = "Times New Roman"
    r_concl.font.size = Pt(11.5)

    doc.save(output_docx)
    print(f"Successfully generated DOCX: {output_docx}")

if __name__ == "__main__":
    generate_outcomes_pdf()
    generate_outcomes_docx()
