"""
Report generation endpoints.
Generates PDF, Excel and CSV files from live database data.
"""
import io
import csv
from datetime import datetime, timedelta

from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Header
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from backend.database import get_db
from backend.models import (
    Vehicle, Driver, TripLog, BatteryLog, SpeedLog,
    RevenueLog, MaintenanceLog, AlertLog, DriverRealtimeData,
)
from backend.auth import get_current_admin_or_driver, verify_token_and_get_user

router = APIRouter(prefix="/api/reports", tags=["Reports"])

def get_auth_user_flexible(
    token: Optional[str] = Query(None),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db)
):
    raw_token = None
    if token:
        raw_token = token
    elif authorization:
        parts = authorization.split(" ")
        if len(parts) == 2 and parts[0].lower() == "bearer":
            raw_token = parts[1]
        else:
            raw_token = authorization
    if not raw_token:
        raise HTTPException(status_code=401, detail="Authentication required")
    return verify_token_and_get_user(raw_token, db)


# ───────────────────────────────────────────────
# Helpers
# ───────────────────────────────────────────────

def _streaming(buffer: io.BytesIO, filename: str, media_type: str):
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


def _clean_pdf_text(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    replacements = {
        "—": "-",
        "–": "-",
        "°": " deg ",
        "₹": "Rs. ",
        "•": "*",
        "…": "...",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "±": "+/-",
    }
    for old, new in replacements.items():
        s = s.replace(old, new)
    # Ensure string encodes cleanly in latin-1 / core fonts
    return s.encode("latin-1", errors="replace").decode("latin-1")


def _build_pdf(title: str, headers: list[str], rows: list[list]) -> io.BytesIO:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 12, _clean_pdf_text(title), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 6, f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)

    # Column widths – distribute evenly across page width
    col_count = len(headers)
    usable_width = pdf.w - pdf.l_margin - pdf.r_margin
    col_w = usable_width / col_count

    # Header row
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(40, 40, 60)
    pdf.set_text_color(255, 255, 255)
    for h in headers:
        pdf.cell(col_w, 8, _clean_pdf_text(h), border=1, fill=True, align="C")
    pdf.ln()

    # Data rows
    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(0, 0, 0)
    for i, row in enumerate(rows):
        if i % 2 == 0:
            pdf.set_fill_color(245, 245, 250)
        else:
            pdf.set_fill_color(255, 255, 255)
        for val in row:
            clean_val = _clean_pdf_text(val)
            pdf.cell(col_w, 7, clean_val[:40], border=1, fill=True, align="C")
        pdf.ln()

    buf = io.BytesIO()
    pdf.output(buf)
    return buf


def _build_excel(title: str, headers: list[str], rows: list[list]) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel sheet name max 31 chars

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center")

    # Timestamp row
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ts_cell = ws.cell(row=2, column=1, value=f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    ts_cell.font = Font(name="Calibri", size=9, italic=True, color="666666")
    ts_cell.alignment = Alignment(horizontal="center")

    # Header row (row 4)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows (row 5+)
    for row_idx, row in enumerate(rows, 5):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal="center")

    # Auto-fit column widths (approximate)
    for col_idx in range(1, len(headers) + 1):
        max_len = max(
            len(str(headers[col_idx - 1])),
            *[len(str(r[col_idx - 1])) for r in rows] if rows else [0],
        )
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A"].width = min(max_len + 4, 30)

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def _build_csv(headers: list[str], rows: list[list]) -> io.BytesIO:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return io.BytesIO(buf.getvalue().encode("utf-8"))


# ───────────────────────────────────────────────
# Data-fetching helpers (one per report type)
# ───────────────────────────────────────────────

def _fleet_performance_data(db: Session):
    vehicles = db.query(Vehicle).all()
    headers = ["Vehicle", "Brand", "Model", "Status", "Battery %", "Health %", "Range (km)", "Trips", "Distance (km)", "Energy (kWh)"]
    rows = []
    for v in vehicles:
        trips_count = db.query(TripLog).filter(TripLog.vehicle_id == v.vehicle_id).count()
        total_dist = db.query(func.sum(TripLog.distance_covered_km)).filter(TripLog.vehicle_id == v.vehicle_id).scalar() or 0.0
        total_energy = db.query(func.sum(TripLog.energy_consumed_kwh)).filter(TripLog.vehicle_id == v.vehicle_id).scalar() or 0.0
        rows.append([
            v.car_name or f"{v.brand} {v.model}",
            v.brand, v.model,
            v.status or "In Garage",
            round(v.current_soc, 1),
            round(v.battery_health, 1),
            round(v.claimed_range_km or 0, 1),
            trips_count,
            round(total_dist, 1),
            round(total_energy, 1),
        ])
    return "Fleet Performance Report", headers, rows


def _driver_performance_data(db: Session):
    drivers = db.query(Driver).filter(Driver.is_active == True).all()
    headers = ["Driver", "Email", "License", "Trips", "Distance (km)", "Energy (kWh)", "Efficiency (km/kWh)", "Overspeeding Events", "Revenue ($)"]
    rows = []
    for d in drivers:
        trips = db.query(TripLog).filter(TripLog.driver_id == d.driver_id).count()
        dist = db.query(func.sum(TripLog.distance_covered_km)).filter(TripLog.driver_id == d.driver_id).scalar() or 0.0
        energy = db.query(func.sum(TripLog.energy_consumed_kwh)).filter(TripLog.driver_id == d.driver_id).scalar() or 0.0
        overspeed = db.query(func.sum(TripLog.overspeeding_events_count)).filter(TripLog.driver_id == d.driver_id).scalar() or 0
        revenue = db.query(func.sum(RevenueLog.amount)).filter(RevenueLog.driver_id == d.driver_id).scalar() or 0.0
        eff = round(dist / energy, 2) if energy > 0 else 0.0
        rows.append([
            f"{d.first_name} {d.last_name}",
            d.email or "-",
            d.license_number,
            trips,
            round(dist, 1),
            round(energy, 1),
            eff,
            int(overspeed),
            round(revenue, 2),
        ])
    return "Driver Performance Report", headers, rows


def _battery_health_data(db: Session):
    vehicles = db.query(Vehicle).all()
    headers = ["Vehicle", "Battery Capacity (kWh)", "Current SoC %", "Health %", "Initial Health %", "Latest Voltage (V)", "Latest Temp (°C)"]
    rows = []
    for v in vehicles:
        last_bat = db.query(BatteryLog).filter(BatteryLog.vehicle_id == v.vehicle_id).order_by(BatteryLog.time.desc()).first()
        voltage = round(last_bat.voltage, 1) if last_bat and last_bat.voltage else "-"
        temp = round(last_bat.temperature_celsius, 1) if last_bat and last_bat.temperature_celsius else "-"
        rows.append([
            v.car_name or f"{v.brand} {v.model}",
            v.battery_capacity_kwh,
            round(v.current_soc, 1),
            round(v.battery_health, 1),
            round(v.initial_battery_health_pct or 100, 1),
            voltage,
            temp,
        ])
    return "Battery Health Report", headers, rows


def _maintenance_data(db: Session):
    records = (
        db.query(MaintenanceLog, Vehicle)
        .join(Vehicle, MaintenanceLog.vehicle_id == Vehicle.vehicle_id)
        .order_by(MaintenanceLog.performed_at.desc())
        .all()
    )
    headers = ["Date", "Vehicle", "Description", "Cost ($)"]
    rows = []
    for m, v in records:
        rows.append([
            m.performed_at.strftime("%Y-%m-%d"),
            v.car_name or f"{v.brand} {v.model}",
            m.description,
            round(m.cost, 2),
        ])
    # If no maintenance records exist, provide a summary row
    if not rows:
        rows.append(["-", "No maintenance records found", "-", 0])
    return "Maintenance Report", headers, rows


def _energy_consumption_data(db: Session):
    three_months_ago = datetime.utcnow() - timedelta(days=90)
    vehicles = db.query(Vehicle).all()
    headers = ["Vehicle", "Total Energy (kWh)", "Est. Cost ($)", "Trips", "Avg kWh/Trip"]
    rows = []
    for v in vehicles:
        energy = db.query(func.sum(TripLog.energy_consumed_kwh)).filter(
            TripLog.vehicle_id == v.vehicle_id, TripLog.time >= three_months_ago
        ).scalar() or 0.0
        trips = db.query(TripLog).filter(
            TripLog.vehicle_id == v.vehicle_id, TripLog.time >= three_months_ago
        ).count()
        avg = round(energy / trips, 2) if trips > 0 else 0.0
        rows.append([
            v.car_name or f"{v.brand} {v.model}",
            round(energy, 1),
            round(energy * 0.14, 2),
            trips,
            avg,
        ])
    return "Energy Consumption Report", headers, rows


def _revenue_analytics_data(db: Session):
    three_months_ago = datetime.utcnow() - timedelta(days=90)
    drivers = db.query(Driver).filter(Driver.is_active == True).all()
    headers = ["Driver", "Trips", "Revenue ($)", "Distance (km)", "Energy Cost ($)", "Profit Margin ($)"]
    rows = []
    for d in drivers:
        trips = db.query(TripLog).filter(TripLog.driver_id == d.driver_id, TripLog.time >= three_months_ago).count()
        rev = db.query(func.sum(RevenueLog.amount)).filter(
            RevenueLog.driver_id == d.driver_id, RevenueLog.logged_at >= three_months_ago
        ).scalar() or 0.0
        dist = db.query(func.sum(TripLog.distance_covered_km)).filter(
            TripLog.driver_id == d.driver_id, TripLog.time >= three_months_ago
        ).scalar() or 0.0
        energy = db.query(func.sum(TripLog.energy_consumed_kwh)).filter(
            TripLog.driver_id == d.driver_id, TripLog.time >= three_months_ago
        ).scalar() or 0.0
        energy_cost = round(energy * 0.14, 2)
        profit = round(rev - energy_cost, 2)
        rows.append([
            f"{d.first_name} {d.last_name}",
            trips,
            round(rev, 2),
            round(dist, 1),
            energy_cost,
            profit,
        ])
    return "Revenue Analytics Report", headers, rows


# Map report slug → data-builder
_REPORT_MAP = {
    "fleet-performance": _fleet_performance_data,
    "driver-performance": _driver_performance_data,
    "battery-health": _battery_health_data,
    "maintenance": _maintenance_data,
    "energy-consumption": _energy_consumption_data,
    "revenue-analytics": _revenue_analytics_data,
}


# ───────────────────────────────────────────────
# Endpoints
# ───────────────────────────────────────────────

import os

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
REPORTS_DIR = os.path.join(PROJECT_ROOT, "Reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


@router.get("/saved")
def get_saved_reports(current_user=Depends(get_current_admin_or_driver)):
    """List all report files currently saved in the project's Reports folder."""
    os.makedirs(REPORTS_DIR, exist_ok=True)
    files = []
    for fname in os.listdir(REPORTS_DIR):
        fpath = os.path.join(REPORTS_DIR, fname)
        if os.path.isfile(fpath):
            stat = os.stat(fpath)
            ext = fname.split(".")[-1].lower() if "." in fname else "file"
            files.append({
                "filename": fname,
                "format": ext,
                "size_bytes": stat.st_size,
                "size_kb": round(stat.st_size / 1024, 1),
                "created_at": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            })
    files.sort(key=lambda x: x["created_at"], reverse=True)
    return files


@router.get("/saved/{filename}")
def download_saved_report(filename: str, current_user=Depends(get_auth_user_flexible)):
    """Download a specific file saved in the Reports folder."""
    safe_name = os.path.basename(filename)
    fpath = os.path.join(REPORTS_DIR, safe_name)
    if not os.path.exists(fpath):
        raise HTTPException(status_code=404, detail="Report file not found")
    
    media_type = "application/pdf" if safe_name.endswith(".pdf") else (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if safe_name.endswith(".xlsx") else "text/csv"
    )
    with open(fpath, "rb") as f:
        content = f.read()
    buf = io.BytesIO(content)
    return _streaming(buf, safe_name, media_type)


@router.get("/download/{report_type}/{file_format}")
def download_report(
    report_type: str,
    file_format: str,
    db: Session = Depends(get_db),
    current_user=Depends(get_auth_user_flexible),
):
    builder = _REPORT_MAP.get(report_type)
    if not builder:
        raise HTTPException(status_code=400, detail=f"Unknown report type: {report_type}")

    if file_format not in ("pdf", "excel", "csv"):
        raise HTTPException(status_code=400, detail=f"Unsupported format: {file_format}")

    title, headers, rows = builder(db)
    safe_name = report_type.replace("-", "_")
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

    if file_format == "pdf":
        buf = _build_pdf(title, headers, rows)
        filename = f"{safe_name}_{ts}.pdf"
        media_type = "application/pdf"
    elif file_format == "excel":
        buf = _build_excel(title, headers, rows)
        filename = f"{safe_name}_{ts}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:  # csv
        buf = _build_csv(headers, rows)
        filename = f"{safe_name}_{ts}.csv"
        media_type = "text/csv"

    # Automatically save a copy into the project's 'Reports' folder
    try:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        disk_path = os.path.join(REPORTS_DIR, filename)
        with open(disk_path, "wb") as f_out:
            f_out.write(buf.getvalue())
        print(f"[REPORTS] Saved {filename} to {disk_path}")
    except Exception as save_err:
        print(f"[REPORTS ERROR] Could not save to disk: {save_err}")

    return _streaming(buf, filename, media_type)

