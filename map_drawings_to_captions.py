import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'Project PPT\Rough Report\VOLTFLEET AI (AutoRecovered).docx')

count = 0
for idx, p in enumerate(doc.paragraphs):
    if "drawing" in p._element.xml:
        count += 1
        before = [doc.paragraphs[i].text.strip() for i in range(max(0, idx-3), idx) if doc.paragraphs[i].text.strip()]
        after = [doc.paragraphs[i].text.strip() for i in range(idx+1, min(len(doc.paragraphs), idx+4)) if doc.paragraphs[i].text.strip()]
        print(f"\n--- Drawing #{count} at Para {idx} ---")
        print("  Before:", " // ".join(before))
        print("  After :", " // ".join(after))
